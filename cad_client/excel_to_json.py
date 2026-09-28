from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET


try:
    from openpyxl import load_workbook
except Exception:  # pragma: no cover
    load_workbook = None


NS = {
    'a': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    'p': 'http://schemas.openxmlformats.org/package/2006/relationships',
}


def normalize_scalar(value):
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        if value == '':
            return ''
        if value.lower() in {'true', 'false'}:
            return value.lower() == 'true'
        if re.fullmatch(r'[-+]?\d+', value):
            return int(value)
        if re.fullmatch(r'[-+]?\d+\.\d+([eE][-+]?\d+)?', value):
            return float(value)
        return value
    return value


def col_index_from_ref(cell_ref: str | None) -> int:
    if not cell_ref:
        return 0
    match = re.match(r'([A-Za-z]+)', cell_ref)
    if not match:
        return 0
    letters = match.group(1).upper()
    idx = 0
    for ch in letters:
        idx = idx * 26 + (ord(ch) - 64)
    return idx - 1


def parse_xlsx_with_stdlib(xlsx_path: Path) -> dict:
    def read_shared_strings(zf: zipfile.ZipFile):
        if 'xl/sharedStrings.xml' not in zf.namelist():
            return []
        root = ET.fromstring(zf.read('xl/sharedStrings.xml'))
        shared = []
        for si in root.findall('.//a:si', NS):
            texts = [t.text or '' for t in si.findall('.//a:t', NS)]
            shared.append(''.join(texts))
        return shared

    def parse_cell(cell, shared_strings):
        cell_type = cell.attrib.get('t')
        value_node = cell.find('a:v', NS)
        if value_node is None:
            if cell_type == 'inlineStr':
                texts = ''.join(t.text or '' for t in cell.findall('.//a:t', NS))
                return normalize_scalar(texts)
            return None

        raw = value_node.text
        if cell_type == 's':
            try:
                idx = int(raw)
            except (TypeError, ValueError):
                return raw
            return normalize_scalar(shared_strings[idx] if 0 <= idx < len(shared_strings) else raw)
        if cell_type == 'b':
            return raw == '1'
        return normalize_scalar(raw)

    with zipfile.ZipFile(xlsx_path) as zf:
        workbook_root = ET.fromstring(zf.read('xl/workbook.xml'))
        rels_root = ET.fromstring(zf.read('xl/_rels/workbook.xml.rels'))
        rel_map = {
            rel.attrib['Id']: rel.attrib['Target']
            for rel in rels_root.findall('p:Relationship', NS)
        }
        shared_strings = read_shared_strings(zf)

        result = {}
        for sheet in workbook_root.findall('a:sheets/a:sheet', NS):
            sheet_name = sheet.attrib.get('name', 'Sheet')
            rid = sheet.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
            target = rel_map.get(rid)
            if not target:
                result[sheet_name] = []
                continue

            sheet_xml = ET.fromstring(zf.read('xl/' + target))
            rows = []
            for row in sheet_xml.findall('.//a:sheetData/a:row', NS):
                row_values = {}
                for cell in row.findall('a:c', NS):
                    ref = cell.attrib.get('r')
                    col_idx = col_index_from_ref(ref)
                    row_values[col_idx] = parse_cell(cell, shared_strings)
                if row_values:
                    max_col = max(row_values) + 1
                    values = [row_values.get(i) for i in range(max_col)]
                    rows.append(values)
                else:
                    rows.append([])

            result[sheet_name] = rows

        return result


def parse_xlsx_with_openpyxl(xlsx_path: Path) -> dict:
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    result = {}
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            values = [normalize_scalar(v) for v in row]
            if any(v is not None and str(v).strip() != '' for v in values):
                rows.append(values)
            else:
                rows.append([])
        result[ws.title] = rows
    return result


def excel_to_json(xlsx_path: Path) -> dict:
    if load_workbook is not None:
        try:
            return parse_xlsx_with_openpyxl(xlsx_path)
        except Exception:
            return parse_xlsx_with_stdlib(xlsx_path)
    return parse_xlsx_with_stdlib(xlsx_path)


def main() -> None:
    base_dir = Path(__file__).resolve().parent
    excel_candidates = sorted(base_dir.glob('*.xlsx'))
    if not excel_candidates:
        raise FileNotFoundError(f'未找到 .xlsx 文件，请将目标 Excel 放到 {base_dir} 目录下。')

    excel_path = excel_candidates[0]
    output_path = excel_path.with_suffix('.json')

    if load_workbook is not None:
        print(f'已检测到第三方库 openpyxl，优先使用 openpyxl 解析。')
    else:
        print('未检测到第三方库 openpyxl，使用 Python 标准库回退方案解析。')

    data = excel_to_json(excel_path)
    with output_path.open('w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

    print(f'转换完成：{excel_path} -> {output_path}')
    print(f'工作表数量：{len(data)}')
    for sheet_name in data:
        print(f'  - {sheet_name}: {len(data[sheet_name])} 条记录')


if __name__ == '__main__':
    main()
